from __future__ import annotations

import io

from openpyxl import Workbook

from boards.macro_national import analytics, parser


def _xlsx_from_rows(headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for j, h in enumerate(headers, start=1):
        ws.cell(1, j, h)
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_economie_budgetaire_scalar_parse() -> None:
    content = _xlsx_from_rows(
        ["fraude", "rescoring", "total"],
        [[1_000_000, 2_500_000, 3_500_000]],
    )
    values = parser.parse_economie_budgetaire(content)
    eb = analytics.compute_economie_budgetaire(values)
    assert eb.fraude == 1_000_000
    assert eb.rescoring == 2_500_000
    assert eb.total == 3_500_000


def test_detect_economie_from_filename() -> None:
    content = _xlsx_from_rows(
        ["fraude", "rescoring", "total"],
        [[100, 200, 300]],
    )
    kind = parser.detect_file_kind("economie_budgetaire_2026.xlsx", content)
    assert kind is not None
    assert kind.value == "economie_budgetaire"


def test_region_bloques_sums_all_categories(consolidated_xlsx_bytes: bytes) -> None:
    df = parser.parse_consolidated(consolidated_xlsx_bytes)
    bloques = analytics.compute_region_bloques(df)
    by_region = {r.region: r.bloques for r in bloques}
    # Béni Mellal-Khénifra: (20+15+30+40+25+50) + (25+18+35+45+28+55) = 180 + 206 = 386
    assert by_region["Béni Mellal-Khénifra"] == 386
    # Casablanca-Settat: (80+60+100) + (85+65+110) = 240 + 260 = 500
    assert by_region["Casablanca-Settat"] == 500


def test_parse_consolidated_rejects_missing_columns() -> None:
    content = _xlsx_from_rows(
        ["month", "region", "province"],
        [["2026-01", "R1", "P1"]],
    )
    try:
        parser.parse_consolidated(content)
        raise AssertionError("should have raised ValueError")
    except ValueError as exc:
        assert "colonnes manquantes" in str(exc)
