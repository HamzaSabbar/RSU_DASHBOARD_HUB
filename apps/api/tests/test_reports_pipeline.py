from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import to_excel
from sqlalchemy import select

from config import settings
from db.models import ReportJob, ReportUploadBatch
from db.session import SessionLocal, engine
from reports import calculations, facts, parser, service, template
from reports.repository import LocalJsonReportJobRepository, ReportJobRecord
from reports.storage import LocalObjectStorage, get_storage


def _add_section(
    ws: Any,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    *,
    omit: set[str] | None = None,
) -> None:
    omit = omit or set()
    if ws.max_row > 1 or ws.cell(1, 1).value is not None:
        ws.append([])
    ws.append([title])
    effective_headers = [header for header in headers if header not in omit]
    ws.append(effective_headers)
    for row in rows:
        row_map = dict(zip(headers, row, strict=True))
        ws.append([row_map.get(header) for header in effective_headers])


def make_report_workbook(
    *,
    missing_sheet: str | None = None,
    omit_rsu_stock_total: bool = False,
    invalid_region: bool = False,
    province_mismatch: bool = False,
    id_chargement: str = "LOAD-001",
    include_id_chargement: bool = True,
) -> bytes:
    wb = Workbook()
    default = wb.active
    assert default is not None
    wb.remove(default)

    ws = wb.create_sheet("01_Parametres")
    parameter_headers = [
        "id_chargement",
        "version_fichier",
        "date_rapport",
        "debut_periode",
        "fin_periode",
        "date_reference_donnees",
        "date_heure_extraction",
        "systeme_source",
        "code_langue",
        "mois_reporting_courant",
        "indicateur_mois_partiel",
        "suffixe_libelle_mois_partiel",
        "type_unite_graphique_rsu",
        "methode_tendance",
        "utiliser_regles_montant_secours",
        "charge_par",
        "commentaires",
    ]
    _add_section(
        ws,
        "Métadonnées",
        parameter_headers,
        [
            [
                id_chargement,
                "1.0",
                to_excel(date(2026, 4, 24)),
                "2026-03-01",
                "2026-04-30",
                "2026-04-20",
                "2026-04-20T08:30:00",
                "RSU",
                "fr",
                "2026-04",
                True,
                "(partiel)",
                "MENAGES",
                "REGRESSION_LINEAIRE",
                True,
                "tests",
                "",
            ]
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )

    ws = wb.create_sheet("02_Regions")
    _add_section(
        ws,
        "Régions",
        ["code_region", "nom_region", "ordre_affichage"],
        [["CAS", "Casablanca-Settat", 1], ["RSK", "Rabat-Salé-Kénitra", 2]],
    )

    ws = wb.create_sheet("03_Provinces")
    _add_section(
        ws,
        "Provinces",
        ["code_region", "nom_province", "ordre_affichage"],
        [["CAS", "CASABLANCA", 1], ["RSK", "KÉNITRA", 2]],
    )

    ws = wb.create_sheet("04_Codes")
    code_rows = []
    for type_code, values in {
        "code_registre": ["RNP", "RSU"],
        "type_unite": ["MENAGES", "PERSONNES"],
        "mode_source": ["SAISIE", "DIFF_SNAPSHOT_CUMULE", "FLUX_PERIODE"],
        "code_langue": ["fr"],
        "methode_tendance": ["REGRESSION_LINEAIRE", "MOYENNE_MOBILE", "AUCUNE"],
        "code_type_famille": ["T1", "T2"],
        "code_niveau_risque": ["ELEVE", "MOYEN"],
        "code_perimetre_programme": ["ASD", "AMO_TADAMON"],
        "code_motif_blocage": ["FMS", "MULTI", "INDIV"],
        "code_programme": ["ASD", "AMO_TADAMON"],
        "unite_beneficiaire": ["MENAGES"],
        "type_montant": ["STANDARD"],
    }.items():
        for index, value in enumerate(values, start=1):
            code_rows.append([type_code, value, value, index])
    _add_section(ws, "Codes", ["type_code", "code", "libelle", "ordre_affichage"], code_rows)

    ws = wb.create_sheet("10_RSU")
    rsu_stock_headers = [
        "id_chargement",
        "date_reference",
        "code_registre",
        "type_unite",
        "total_cumule",
        "systeme_source",
        "commentaires",
    ]
    _add_section(
        ws,
        "Stock RNP/RSU",
        rsu_stock_headers,
        [
            [id_chargement, "2026-04-20", "RNP", "PERSONNES", 12_000_000, "RNP", ""],
            [id_chargement, "2026-04-20", "RSU", "MENAGES", 4_000_000, "RSU", ""],
            [id_chargement, "2026-04-20", "RSU", "PERSONNES", 11_000_000, "RSU", ""],
        ],
        omit=(
            ({"total_cumule"} if omit_rsu_stock_total else set())
            | (set() if include_id_chargement else {"id_chargement"})
        ),
    )
    flow_region = "XYZ" if invalid_region else "CAS"
    mismatch_region = "CAS" if province_mismatch else "RSK"
    _add_section(
        ws,
        "Nouvelles inscriptions",
        [
            "id_chargement",
            "debut_periode",
            "fin_periode",
            "date_evenement",
            "mois_evenement",
            "code_region",
            "nom_province",
            "code_registre",
            "type_unite",
            "nb_nouvelles_inscriptions",
            "mode_source",
            "systeme_source",
            "commentaires",
        ],
        [
            [id_chargement, "2026-03-01", "2026-03-31", "2026-03-10", "2026-03", "CAS", "CASABLANCA", "RNP", "PERSONNES", 100, "FLUX_PERIODE", "RNP", ""],
            [id_chargement, "2026-03-01", "2026-03-31", "2026-03-11", "2026-03", "RSK", "KÉNITRA", "RNP", "PERSONNES", 80, "FLUX_PERIODE", "RNP", ""],
            [id_chargement, "2026-04-01", "2026-04-30", "2026-04-10", "2026-04", "CAS", "CASABLANCA", "RNP", "PERSONNES", 150, "FLUX_PERIODE", "RNP", ""],
            [id_chargement, "2026-04-01", "2026-04-30", "2026-04-11", "2026-04", "RSK", "KÉNITRA", "RNP", "PERSONNES", 130, "FLUX_PERIODE", "RNP", ""],
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )
    _add_section(
        ws,
        "Annotations",
        ["code_graphique", "mois_evenement", "libelle_annotation", "commentaires"],
        [["RSU", "2026-04", "Mois partiel selon extraction source.", ""]],
    )

    ws = wb.create_sheet("12_RSU_Nouvelles_Inscriptions")
    _add_section(
        ws,
        "Nouvelles inscriptions RSU ménages",
        [
            "id_chargement",
            "debut_periode",
            "fin_periode",
            "date_evenement",
            "mois_evenement",
            "code_region",
            "nom_province",
            "nb_nouveaux_menages_rsu",
            "mode_source",
            "systeme_source",
            "commentaires",
        ],
        [
            [id_chargement, "2026-03-01", "2026-03-31", "2026-03-10", "2026-03", "CAS", "CASABLANCA", 50, "FLUX_PERIODE", "RSU", ""],
            [id_chargement, "2026-03-01", "2026-03-31", "2026-03-11", "2026-03", "RSK", "KÉNITRA", 30, "FLUX_PERIODE", "RSU", ""],
            [id_chargement, "2026-04-01", "2026-04-30", "2026-04-10", "2026-04", "CAS", "CASABLANCA", 70, "FLUX_PERIODE", "RSU", ""],
            [id_chargement, "2026-04-01", "2026-04-30", "2026-04-11", "2026-04", "RSK", "KÉNITRA", 40, "FLUX_PERIODE", "RSU", ""],
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )

    _program_sheet(
        wb,
        "20_ASD",
        id_chargement,
        flow_region=flow_region,
        mismatch_region=mismatch_region,
        programme="ASD",
        missing_rescoring_amount=False,
        missing_fraud_amount=False,
        include_id_chargement=include_id_chargement,
    )
    _program_sheet(
        wb,
        "30_AMO_Tadamon",
        id_chargement,
        flow_region="CAS",
        mismatch_region="RSK",
        programme="AMO",
        missing_rescoring_amount=True,
        missing_fraud_amount=True,
        include_id_chargement=include_id_chargement,
    )
    _fms_sheet(wb, id_chargement, include_id_chargement=include_id_chargement)
    _amount_rules_sheet(wb)

    if missing_sheet:
        del wb[missing_sheet]

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_split_table_workbook(
    *,
    include_system_sheets: bool = True,
    include_rsu_annotations: bool = True,
    minimal_parameters: bool = False,
    client_data_only: bool = False,
) -> bytes:
    parsed = parser.parse_workbook(make_report_workbook(include_id_chargement=False))
    assert parsed.validation.summary.errors == 0

    wb = Workbook()
    default = wb.active
    assert default is not None
    wb.remove(default)

    system_sheet_keys = {"regions", "provinces", "codes", "amount_rules"}
    minimal_parameter_headers = {
        "debut_periode",
        "fin_periode",
    }
    client_omitted_headers = {
        "date_reference",
        "date_evenement",
        "mois_evenement",
        "code_region",
        "code_registre",
        "code_perimetre_programme",
        "mode_source",
        "systeme_source",
        "commentaires",
    }

    for spec in parser.SECTION_SPECS:
        if not include_system_sheets and spec.key in system_sheet_keys:
            continue
        if not include_rsu_annotations and spec.key == "rsu_annotations":
            continue
        if client_data_only and spec.key in {"rsu_stock", "asd_stock", "amo_stock"}:
            continue
        if spec.key == "parameters":
            rows = [parsed.normalized["metadata"]]
        else:
            rows = parsed.normalized.get(spec.key, [])
        if not rows and not spec.required:
            continue

        sheet_name = spec.sheet_aliases[0] if spec.sheet_aliases else spec.sheet
        headers = [
            column
            for column in spec.columns
            if column != "id_chargement"
            and (not minimal_parameters or spec.key != "parameters" or column in minimal_parameter_headers)
            and (
                not client_data_only
                or spec.key == "parameters"
                or column not in client_omitted_headers
            )
            and (not client_data_only or spec.key != "parameters" or column in minimal_parameter_headers)
            and (not client_data_only or spec.key != "rsu_annotations")
        ]
        ws = wb.create_sheet(sheet_name)
        client_headers = []
        for header in headers:
            client_header = header
            if client_data_only:
                client_header = {
                    "code_type_famille": "type_famille",
                    "code_niveau_risque": "niveau_risque",
                    "code_motif_blocage": "motif_blocage",
                }.get(header, header)
                if spec.key == "rsu_household_registrations" and header == "nb_nouvelles_inscriptions":
                    client_header = "nb_nouveaux_menages_rsu"
            client_headers.append(client_header)
        ws.append(client_headers)
        for row in rows:
            values = []
            for header in headers:
                value = row.get(header)
                if client_data_only and header in {"nom_region", "nom_province"} and value:
                    value = str(value).lower()
                values.append(value)
            ws.append(values)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _program_sheet(
    wb: Workbook,
    sheet_name: str,
    id_chargement: str,
    *,
    flow_region: str,
    mismatch_region: str,
    programme: str,
    missing_rescoring_amount: bool,
    missing_fraud_amount: bool,
    include_id_chargement: bool,
) -> None:
    ws = wb.create_sheet(sheet_name)
    _add_section(
        ws,
        "Stock",
        ["id_chargement", "date_reference", "type_unite", "nb_actifs", "systeme_source", "commentaires"],
        [
            [id_chargement, "2026-04-20", "MENAGES", 1000 if programme == "ASD" else 700, programme, ""],
            [id_chargement, "2026-04-20", "PERSONNES", 3300 if programme == "ASD" else 2100, programme, ""],
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )
    _add_section(
        ws,
        "Flux",
        [
            "id_chargement",
            "debut_periode",
            "fin_periode",
            "date_evenement",
            "mois_evenement",
            "code_region",
            "nom_province",
            "nb_entrants_menages",
            "nb_sortants_menages",
            "nb_entrants_personnes",
            "nb_sortants_personnes",
            "montant_mensuel_entrants_dh",
            "montant_mensuel_sortants_dh",
            "systeme_source",
            "commentaires",
        ],
        [
            [id_chargement, "2026-03-01", "2026-03-31", "2026-03-05", "2026-03", "CAS", "CASABLANCA", 40, 30, 120, 90, 4000, 3000, programme, ""],
            [id_chargement, "2026-03-01", "2026-03-31", "2026-03-06", "2026-03", "RSK", "KÉNITRA", 20, 15, 60, 45, 2000, 1500, programme, ""],
            [id_chargement, "2026-04-01", "2026-04-30", "2026-04-05", "2026-04", flow_region, "CASABLANCA", 50 if programme == "ASD" else 20, 20 if programme == "ASD" else 5, 150, 60, 5000, 2000, programme, ""],
            [id_chargement, "2026-04-01", "2026-04-30", "2026-04-06", "2026-04", mismatch_region, "KÉNITRA", 30 if programme == "ASD" else 10, 10 if programme == "ASD" else 2, 90, 30, 3000, 1000, programme, ""],
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )
    _add_section(
        ws,
        "Rescoring",
        [
            "id_chargement",
            "debut_periode",
            "fin_periode",
            "date_evenement",
            "mois_evenement",
            "code_region",
            "nom_province",
            "nb_sortants_menages",
            "nb_sortants_personnes",
            "nb_menages_non_communiques",
            "montant_mensuel_arrete_dh",
            "systeme_source",
            "commentaires",
        ],
        [
            [
                id_chargement,
                "2026-04-01",
                "2026-04-30",
                "2026-04-07",
                "2026-04",
                "CAS" if programme == "ASD" else "RSK",
                "CASABLANCA" if programme == "ASD" else "KÉNITRA",
                5 if programme == "ASD" else 3,
                15 if programme == "ASD" else 9,
                2 if programme == "ASD" else 1,
                1000 if not missing_rescoring_amount else None,
                programme,
                "",
            ]
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )
    _add_section(
        ws,
        "Fraude",
        [
            "id_chargement",
            "debut_periode",
            "fin_periode",
            "date_evenement",
            "mois_evenement",
            "code_region",
            "nom_province",
            "nb_menages_radies",
            "nb_personnes_radiees",
            "montant_mensuel_arrete_dh",
            "systeme_source",
            "commentaires",
        ],
        [
            [
                id_chargement,
                "2026-04-01",
                "2026-04-30",
                "2026-04-08",
                "2026-04",
                "CAS" if programme == "ASD" else "RSK",
                "CASABLANCA" if programme == "ASD" else "KÉNITRA",
                2 if programme == "ASD" else 4,
                6 if programme == "ASD" else 12,
                5000 if not missing_fraud_amount else None,
                programme,
                "",
            ]
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )


def _fms_sheet(
    wb: Workbook,
    id_chargement: str,
    *,
    include_id_chargement: bool,
) -> None:
    ws = wb.create_sheet("40_FMS")
    _add_section(
        ws,
        "Traitement FMS",
        [
            "id_chargement",
            "date_reference",
            "debut_periode",
            "fin_periode",
            "date_evenement",
            "mois_evenement",
            "code_type_famille",
            "code_niveau_risque",
            "code_perimetre_programme",
            "code_region",
            "nom_province",
            "demandes_injectees",
            "demandes_traitees",
            "doute_confirme",
            "doute_leve",
            "en_attente",
            "systeme_source",
            "commentaires",
        ],
        [
            [id_chargement, "2026-04-20", "2026-04-01", "2026-04-30", "2026-04-20", "2026-04", "T1", "ELEVE", "ASD", "CAS", "CASABLANCA", 100, 80, 20, 50, 20, "FMS", ""],
            [id_chargement, "2026-04-20", "2026-04-01", "2026-04-30", "2026-04-20", "2026-04", "T2", "MOYEN", "AMO_TADAMON", "RSK", "KÉNITRA", 50, 50, 5, 40, 0, "FMS", ""],
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )
    _add_section(
        ws,
        "Ménages bloqués",
        [
            "id_chargement",
            "date_reference",
            "debut_periode",
            "fin_periode",
            "date_evenement",
            "mois_evenement",
            "code_motif_blocage",
            "code_type_famille",
            "code_region",
            "nom_province",
            "nb_menages_bloques",
            "nb_personnes_bloquees",
            "systeme_source",
            "commentaires",
        ],
        [
            [id_chargement, "2026-04-20", "2026-04-01", "2026-04-30", "2026-04-20", "2026-04", "FMS", "T1", "CAS", "CASABLANCA", 10, 30, "FMS", ""],
            [id_chargement, "2026-04-20", "2026-04-01", "2026-04-30", "2026-04-20", "2026-04", "INDIV", "T1", "CAS", "CASABLANCA", 3, 9, "FMS", ""],
            [id_chargement, "2026-04-20", "2026-04-01", "2026-04-30", "2026-04-20", "2026-04", "MULTI", "T2", "RSK", "KÉNITRA", 5, 15, "FMS", ""],
        ],
        omit=set() if include_id_chargement else {"id_chargement"},
    )


def _amount_rules_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("50_Regles_Montant")
    _add_section(
        ws,
        "Règles montants",
        [
            "code_programme",
            "date_effet_debut",
            "date_effet_fin",
            "code_type_famille",
            "unite_beneficiaire",
            "type_montant",
            "montant_mensuel_dh",
            "facteur_annualisation",
            "commentaires",
        ],
        [
            ["ASD", "2026-01-01", None, "T1", "MENAGES", "STANDARD", 300, 12, ""],
            ["AMO_TADAMON", "2026-01-01", None, "T1", "MENAGES", "STANDARD", 200, 12, ""],
        ],
    )


def _dashboard_from_workbook(content: bytes) -> dict[str, Any]:
    parsed = parser.parse_workbook(content, job_id="job-test")
    assert parsed.validation.summary.errors == 0, [
        message.message for message in parsed.validation.messages
    ]
    return calculations.build_dashboard(
        parsed.normalized,
        job_id="job-test",
        validation=parsed.validation,
    )


def _move_normalized_to_may(normalized: dict[str, Any]) -> dict[str, Any]:
    normalized = {
        key: [dict(row) for row in value] if isinstance(value, list) else dict(value)
        for key, value in normalized.items()
    }
    metadata = normalized["metadata"]
    metadata.update(
        {
            "date_rapport": "2026-05-24",
            "debut_periode": "2026-05-01",
            "fin_periode": "2026-05-31",
            "date_reference_donnees": "2026-05-20",
            "mois_reporting_courant": "2026-05",
        }
    )
    for key in ("rsu_stock", "asd_stock", "amo_stock", "fms_treatment", "fms_blocked"):
        for row in normalized[key]:
            row["date_reference"] = "2026-05-20"
    period_fact_keys = (
        "rsu_new_registrations",
        "rsu_household_registrations",
        "asd_flow",
        "amo_flow",
        "asd_rescoring",
        "amo_rescoring",
        "asd_fraud",
        "amo_fraud",
        "fms_treatment",
        "fms_blocked",
    )
    for key in period_fact_keys:
        normalized[key] = [
            row for row in normalized[key] if row.get("mois_evenement") in {None, "2026-04"}
        ]
    for key in period_fact_keys:
        for row in normalized[key]:
            row["debut_periode"] = "2026-05-01"
            row["fin_periode"] = "2026-05-31"
            row["date_evenement"] = "2026-05-10"
            row["mois_evenement"] = "2026-05"
    for row in normalized["rsu_stock"]:
        if row["code_registre"] == "RSU" and row["type_unite"] == "MENAGES":
            row["total_cumule"] = 4_500_000
    return normalized


def _job_record(job_id: str, *, id_chargement: str, replace: bool = False) -> ReportJobRecord:
    now = datetime.now(tz=UTC)
    return ReportJobRecord(
        id=job_id,
        id_chargement=id_chargement,
        original_filename="rapport.xlsx",
        replace_requested=replace,
        status="running",
        progress=90,
        raw_object_path=f"rsu-dashboard/uploads/{job_id}/input.xlsx",
        normalized_object_path=f"rsu-dashboard/jobs/{job_id}/normalized.json",
        validation_object_path=f"rsu-dashboard/jobs/{job_id}/validation.json",
        dashboard_object_path=f"rsu-dashboard/jobs/{job_id}/dashboard.json",
        created_by=None,
        created_at=now,
        started_at=now,
    )


async def _add_job_row(session: Any, record: ReportJobRecord) -> None:
    session.add(
        ReportJob(
            id=uuid.UUID(record.id),
            id_chargement=record.id_chargement,
            original_filename=record.original_filename,
            replace_requested=record.replace_requested,
            status=record.status,
            progress=record.progress,
            raw_object_path=record.raw_object_path,
            normalized_object_path=record.normalized_object_path,
            validation_object_path=record.validation_object_path,
            dashboard_object_path=record.dashboard_object_path,
            created_by=None,
        )
    )
    await session.flush()


def test_successful_workbook_validation_and_dashboard() -> None:
    parsed = parser.parse_workbook(make_report_workbook(), job_id="job-test")
    assert parsed.validation.summary.errors == 0
    dashboard = calculations.build_dashboard(
        parsed.normalized,
        job_id="job-test",
        validation=parsed.validation,
    )
    assert dashboard["meta"]["idChargement"] == "LOAD-001"
    assert dashboard["cards"]["inscriptions"]["rnpPersonnesTotal"]["raw"] == 460
    assert dashboard["cards"]["inscriptions"]["rsuMenagesTotal"]["raw"] == 190
    assert dashboard["cards"]["inscriptions"]["rsuPersonnesCouvertes"]["raw"] == 11_000_000
    assert "rsuPersonnesTotal" not in dashboard["cards"]["inscriptions"]
    assert dashboard["cards"]["traitementFms"]["demandesTraitees"]["raw"] == 130
    assert dashboard["cards"]["radiationFraude"]["savingAnnualRaw"] == 69_600
    assert dashboard["cards"]["rescoring"]["savingAnnualRaw"] == 19_200
    assert dashboard["cards"]["economieBudgetaire"]["total"]["raw"] == 88_800
    assert "CNRA" in " ".join(dashboard["footnotes"])
    assert "CNSS" in " ".join(dashboard["footnotes"])


def test_workbook_without_id_chargement_is_valid() -> None:
    parsed = parser.parse_workbook(
        make_report_workbook(include_id_chargement=False),
        job_id="job-test",
    )
    assert parsed.validation.summary.errors == 0
    assert not parsed.normalized["metadata"].get("id_chargement")


def test_split_table_workbook_without_section_titles_is_supported() -> None:
    parsed = parser.parse_workbook(make_split_table_workbook(), job_id="job-test")
    assert parsed.validation.summary.errors == 0
    assert parsed.normalized["rsu_stock"][0]["total_cumule"] == 12_000_000
    assert parsed.normalized["asd_flow"][0]["nb_entrants_menages"] == 40


def test_client_fact_workbook_without_system_reference_sheets_is_supported() -> None:
    parsed = parser.parse_workbook(
        make_split_table_workbook(
            include_system_sheets=False,
            include_rsu_annotations=False,
            minimal_parameters=True,
        ),
        job_id="job-test",
    )
    assert parsed.validation.summary.errors == 0
    assert parsed.normalized["regions"][0]["code_region"] == "CAS"
    assert parsed.normalized["provinces"][0]["nom_province"] == "AIN CHOCK"
    assert parsed.normalized["codes"] == []
    assert parsed.normalized["amount_rules"] == []
    assert parsed.normalized["metadata"]["code_langue"] == "fr"
    assert parsed.normalized["metadata"]["mois_reporting_courant"] == "2026-04"


def test_client_data_only_workbook_derives_dates_and_region_codes() -> None:
    parsed = parser.parse_workbook(
        make_split_table_workbook(
            include_system_sheets=False,
            include_rsu_annotations=False,
            minimal_parameters=True,
            client_data_only=True,
        ),
        job_id="job-test",
    )
    assert parsed.validation.summary.errors == 0
    row = parsed.normalized["rsu_new_registrations"][0]
    assert row["debut_periode"] == "2026-03-01"
    assert row["date_evenement"] == row["fin_periode"] == "2026-03-31"
    assert row["mois_evenement"] == "2026-03"
    assert row["code_region"] == "CAS"
    assert row["nom_region"] == "Casablanca-Settat"
    assert row["nom_province"] == "CASABLANCA"
    fms_row = parsed.normalized["fms_treatment"][0]
    assert fms_row["code_type_famille"] == "T1"
    assert fms_row["code_niveau_risque"] == "ELEVE"
    assert fms_row["code_perimetre_programme"] is None
    blocked_row = parsed.normalized["fms_blocked"][0]
    assert blocked_row["code_motif_blocage"] == "FMS"
    assert blocked_row["code_type_famille"] == "T1"


def test_client_data_only_workbook_derives_rnp_rsu_units() -> None:
    workbook = load_workbook(
        io.BytesIO(
            make_split_table_workbook(
                include_system_sheets=False,
                include_rsu_annotations=False,
                minimal_parameters=True,
                client_data_only=True,
            )
        )
    )
    for sheet_name in ("11_RNP_Nouvelles_Inscriptions", "12_RSU_Nouvelles_Inscriptions"):
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        unit_column = headers.index("type_unite") + 1
        worksheet.delete_cols(unit_column)

    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parser.parse_workbook(buffer.getvalue(), job_id="job-test")

    assert parsed.validation.summary.errors == 0
    rnp_row = parsed.normalized["rsu_new_registrations"][0]
    rsu_row = parsed.normalized["rsu_household_registrations"][0]
    assert rnp_row["code_registre"] == "RNP"
    assert rnp_row["type_unite"] == "PERSONNES"
    assert rsu_row["code_registre"] == "RSU"
    assert rsu_row["type_unite"] == "MENAGES"


def test_client_data_only_workbook_without_stock_sheets_is_supported() -> None:
    workbook = load_workbook(
        io.BytesIO(
            make_split_table_workbook(
                include_system_sheets=False,
                include_rsu_annotations=False,
                minimal_parameters=True,
                client_data_only=True,
            )
        )
    )
    for sheet_name in ("10_RSU_Stock", "20_ASD_Stock", "30_AMO_Tadamon_Stock"):
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
    for sheet_name in ("21_ASD_Flux", "31_AMO_Tadamon_Flux"):
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in worksheet[1]]
        for column in (
            "montant_mensuel_sortants_dh",
            "montant_mensuel_entrants_dh",
            "nb_sortants_personnes",
            "nb_entrants_personnes",
        ):
            worksheet.delete_cols(headers.index(column) + 1)
            headers.remove(column)
    worksheet = workbook["41_FMS_Menages_Bloques"]
    headers = [cell.value for cell in worksheet[1]]
    worksheet.delete_cols(headers.index("nb_personnes_bloquees") + 1)

    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parser.parse_workbook(buffer.getvalue(), job_id="job-test")

    assert parsed.validation.summary.errors == 0
    dashboard = calculations.build_dashboard(
        parsed.normalized,
        job_id="job-test",
        validation=parsed.validation,
    )
    assert dashboard["cards"]["inscriptions"]["rnpPersonnesTotal"]["raw"] == 460
    assert dashboard["cards"]["inscriptions"]["rsuMenagesTotal"]["raw"] == 190
    assert dashboard["cards"]["programmesSociaux"]["asdMenagesActifs"]["raw"] == 65
    assert dashboard["cards"]["programmesSociaux"]["amoMenagesActifs"]["raw"] == 38
    assert "asdPersonnesActives" not in dashboard["cards"]["programmesSociaux"]
    assert "amoPersonnesActives" not in dashboard["cards"]["programmesSociaux"]


def test_daily_client_workbook_uses_event_dates_and_person_columns() -> None:
    workbook = load_workbook(
        io.BytesIO(
            make_split_table_workbook(
                include_system_sheets=False,
                include_rsu_annotations=False,
                minimal_parameters=True,
                client_data_only=True,
            )
        )
    )
    daily_values = {
        "11_RNP_Nouvelles_Inscriptions": [
            "2026-03-10",
            "2026-03-11",
            "2026-04-10",
            "2026-04-11",
        ],
        "12_RSU_Nouvelles_Inscriptions": [
            "2026-03-10",
            "2026-03-11",
            "2026-04-10",
            "2026-04-11",
        ],
        "21_ASD_Flux": ["2026-03-05", "2026-03-06", "2026-04-05", "2026-04-06"],
        "22_ASD_Rescoring": ["2026-04-07"],
        "23_ASD_Fraude": ["2026-04-08"],
        "31_AMO_Tadamon_Flux": [
            "2026-03-05",
            "2026-03-06",
            "2026-04-05",
            "2026-04-06",
        ],
        "32_AMO_Tadamon_Rescoring": ["2026-04-07"],
        "33_AMO_Tadamon_Fraude": ["2026-04-08"],
        "40_FMS_Traitement": ["2026-04-20", "2026-04-20"],
        "41_FMS_Menages_Bloques": ["2026-04-20", "2026-04-20", "2026-04-20"],
    }
    for sheet_name, dates in daily_values.items():
        _replace_period_columns_with_event_date(workbook[sheet_name], dates)

    rsu_sheet = workbook["12_RSU_Nouvelles_Inscriptions"]
    headers = [cell.value for cell in rsu_sheet[1]]
    person_col = headers.index("nb_nouvelles_personnes_rsu") + 1
    for row_index, value in enumerate([150, 90, 210, 120], start=2):
        rsu_sheet.cell(row_index, person_col).value = value

    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parser.parse_workbook(buffer.getvalue(), job_id="job-test")

    assert parsed.validation.summary.errors == 0
    first_rnp = parsed.normalized["rsu_new_registrations"][0]
    assert first_rnp["debut_periode"] == "2026-03-10"
    assert first_rnp["fin_periode"] == "2026-03-10"
    assert first_rnp["date_evenement"] == "2026-03-10"
    rsu_person_rows = [
        row
        for row in parsed.normalized["rsu_household_registrations"]
        if row["code_registre"] == "RSU" and row["type_unite"] == "PERSONNES"
    ]
    assert sum(row["nb_nouvelles_inscriptions"] for row in rsu_person_rows) == 570

    dashboard = calculations.build_dashboard(
        parsed.normalized,
        job_id="job-test",
        validation=parsed.validation,
    )
    assert dashboard["cards"]["inscriptions"]["rsuMenagesTotal"]["raw"] == 190
    assert dashboard["cards"]["inscriptions"]["rsuPersonnesCouvertes"]["raw"] == 570


def test_missing_sheet_validation_message_is_french() -> None:
    parsed = parser.parse_workbook(make_report_workbook(missing_sheet="10_RSU"))
    assert parsed.validation.summary.errors > 0
    assert any(
        message.message
        == (
            "Feuille ou section obligatoire manquante: Nouvelles inscriptions RNP "
            "(11_RNP_Nouvelles_Inscriptions ou 11_RSU_Nouvelles_Inscriptions ou 10_RSU)"
        )
        for message in parsed.validation.messages
    )


def test_missing_column_validation_message_is_french() -> None:
    parsed = parser.parse_workbook(make_report_workbook(omit_rsu_stock_total=True))
    assert any(
        message.message
        == "Colonne obligatoire manquante dans la section Stock RNP/RSU: total_cumule"
        for message in parsed.validation.messages
    )


def test_invalid_region_code_validation() -> None:
    parsed = parser.parse_workbook(make_report_workbook(invalid_region=True))
    assert any(message.message == "Code région invalide: XYZ" for message in parsed.validation.messages)


def test_province_region_mismatch_validation() -> None:
    parsed = parser.parse_workbook(make_report_workbook(province_mismatch=True))
    assert any(
        message.message == "La province KÉNITRA n’est pas rattachée à la région CAS"
        for message in parsed.validation.messages
    )


def test_excel_serial_date_parsing() -> None:
    parsed = parser.parse_workbook(make_report_workbook())
    assert parsed.normalized["metadata"]["date_rapport"] == "2026-04-24"


def test_client_template_reference_headers_are_supported() -> None:
    wb = load_workbook(io.BytesIO(make_report_workbook()))

    provinces = wb["03_Provinces"]
    province_header = _find_header_row(provinces, "code_region")
    provinces.cell(province_header, 1).value = "cle_province"
    provinces.cell(province_header, 2).value = "nom_province"
    provinces.cell(province_header, 3).value = "code_region"
    provinces.cell(province_header, 4).value = "nom_region"
    for row_index in range(province_header + 1, provinces.max_row + 1):
        code_region = provinces.cell(row_index, 1).value
        province = provinces.cell(row_index, 2).value
        provinces.cell(row_index, 1).value = str(province).replace(" ", "_")
        provinces.cell(row_index, 2).value = province
        provinces.cell(row_index, 3).value = code_region
        provinces.cell(row_index, 4).value = "Région"

    codes = wb["04_Codes"]
    code_header = _find_header_row(codes, "type_code")
    codes.cell(code_header, 1).value = "nom_champ"
    codes.cell(code_header, 2).value = "code"
    codes.cell(code_header, 3).value = "description"
    codes.cell(code_header, 4).value = None

    buf = io.BytesIO()
    wb.save(buf)
    parsed = parser.parse_workbook(buf.getvalue())
    assert parsed.validation.summary.errors == 0


def test_excel_template_contains_required_sheets_and_headers() -> None:
    wb = load_workbook(io.BytesIO(template.build_excel_template()))
    expected = {
        "00_Guide",
        "01_Parametres",
        "11_RNP_Nouvelles_Inscriptions",
        "12_RSU_Nouvelles_Inscriptions",
        "21_ASD_Flux",
        "22_ASD_Rescoring",
        "23_ASD_Fraude",
        "31_AMO_Tadamon_Flux",
        "32_AMO_Tadamon_Rescoring",
        "33_AMO_Tadamon_Fraude",
        "40_FMS_Traitement",
        "41_FMS_Menages_Bloques",
    }
    assert set(wb.sheetnames) == expected
    assert [cell.value for cell in wb["01_Parametres"][1]][:2] == [
        "debut_periode",
        "fin_periode",
    ]
    assert "nb_nouvelles_inscriptions" in [
        cell.value for cell in wb["11_RNP_Nouvelles_Inscriptions"][1]
    ]
    assert "type_unite" not in [cell.value for cell in wb["11_RNP_Nouvelles_Inscriptions"][1]]
    assert wb["01_Parametres"]["A1"].comment is not None
    assert [cell.value for cell in wb["12_RSU_Nouvelles_Inscriptions"][1]][:5] == [
        "date_evenement",
        "nom_region",
        "nom_province",
        "nb_nouveaux_menages_rsu",
        "nb_nouvelles_personnes_rsu",
    ]
    assert "nb_entrants_personnes" in [cell.value for cell in wb["21_ASD_Flux"][1]]


def _find_header_row(ws: Any, marker: str) -> int:
    for row in ws.iter_rows():
        if any(cell.value == marker for cell in row):
            return int(row[0].row)
    raise AssertionError(f"header {marker!r} not found")


def _replace_period_columns_with_event_date(ws: Any, dates: list[str]) -> None:
    headers = [cell.value for cell in ws[1]]
    if "debut_periode" in headers:
        ws.cell(1, headers.index("debut_periode") + 1).value = "date_evenement"
    elif "date_evenement" not in headers:
        ws.insert_cols(1)
        ws.cell(1, 1).value = "date_evenement"

    headers = [cell.value for cell in ws[1]]
    if "fin_periode" in headers:
        ws.delete_cols(headers.index("fin_periode") + 1)

    headers = [cell.value for cell in ws[1]]
    date_col = headers.index("date_evenement") + 1
    for row_index, value in enumerate(dates, start=2):
        ws.cell(row_index, date_col).value = value


def test_rsu_monthly_evolution_and_linear_trend() -> None:
    dashboard = _dashboard_from_workbook(make_report_workbook())
    values = dashboard["charts"]["rsuInscriptionsMensuelles"]["series"]
    assert [point["raw"] for point in values] == [80, 110]
    assert dashboard["tables"]["rsuEvolution"][1]["deltaRaw"] == 30
    assert calculations.linear_regression_trend([10, 20, 30]) == [10.0, 20.0, 30.0]


def test_asd_net_and_combined_regional_aggregation() -> None:
    dashboard = _dashboard_from_workbook(make_report_workbook())
    asd_current = dashboard["tables"]["asdEvolution"][1]
    assert asd_current["entrantsRaw"] == 80
    assert asd_current["sortantsRaw"] == 30
    assert asd_current["netRaw"] == 50
    regions = {
        row["codeRegion"]: row
        for row in dashboard["charts"]["fluxRegionauxAsdAmot"]["rows"]
    }
    assert regions["CAS"]["entrantsRaw"] == 150
    assert regions["CAS"]["sortantsRaw"] == 85
    assert regions["RSK"]["entrantsRaw"] == 80
    assert dashboard["tables"]["topProvincesFlux"][0]["province"] == "CASABLANCA"


def test_fms_matrix_blocked_and_top_provinces() -> None:
    dashboard = _dashboard_from_workbook(make_report_workbook())
    traitement = dashboard["cards"]["traitementFms"]
    assert traitement["tauxTraitementPct"]["raw"] == pytest.approx(130 / 150)
    matrix = dashboard["charts"]["fmsMatriceRisque"]["rows"]
    assert len(matrix) == 2
    national = dashboard["charts"]["menagesBloquesNational"]["rows"]
    assert sum(row["menagesRaw"] for row in national) == 18
    assert dashboard["tables"]["topProvincesBloquees"][0]["province"] == "CASABLANCA"


@pytest.mark.asyncio
async def test_cumulative_fact_dashboard_filters_by_date_range() -> None:
    async with SessionLocal() as session:
        tx = await session.begin()
        try:
            first = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-CUM-1"))
            second = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-CUM-2"))
            assert first.validation.summary.errors == 0
            assert second.validation.summary.errors == 0

            first_record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-CUM-1")
            second_record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-CUM-2")
            await _add_job_row(session, first_record)
            await _add_job_row(session, second_record)
            await facts.ingest_report_batch(
                session,
                record=first_record,
                normalized=first.normalized,
            )
            await facts.ingest_report_batch(
                session,
                record=second_record,
                normalized=_move_normalized_to_may(second.normalized),
            )

            april = await facts.build_dashboard_for_range(
                session,
                start_date=date(2026, 4, 1),
                end_date=date(2026, 4, 30),
            )
            april_asd = next(
                row for row in april["tables"]["asdEvolution"] if row["month"] == "2026-04"
            )
            assert april["cards"]["inscriptions"]["rsuMenagesTotal"]["raw"] == 110
            assert april_asd["entrantsRaw"] == 80
            assert april_asd["sortantsRaw"] == 30

            april_day = await facts.build_dashboard_for_range(
                session,
                start_date=date(2026, 4, 5),
                end_date=date(2026, 4, 5),
            )
            april_day_asd = next(
                row for row in april_day["tables"]["asdEvolution"] if row["month"] == "2026-04"
            )
            assert april_day["meta"]["dateRange"] == {
                "startDate": "2026-04-05",
                "endDate": "2026-04-05",
            }
            assert april_day_asd["entrantsRaw"] == 50

            april_may = await facts.build_dashboard_for_range(
                session,
                start_date=date(2026, 4, 1),
                end_date=date(2026, 5, 31),
            )
            assert april_may["cards"]["inscriptions"]["rsuMenagesTotal"]["raw"] == 220
            rsu_series = april_may["charts"]["rsuInscriptionsMensuelles"]["series"]
            assert [point["month"] for point in rsu_series] == ["2026-04", "2026-05"]
        finally:
            await tx.rollback()
            await engine.dispose()


@pytest.mark.asyncio
async def test_cumulative_fact_dashboard_rejects_reversed_date_range() -> None:
    async with SessionLocal() as session:
        tx = await session.begin()
        try:
            parsed = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-REVERSED"))
            assert parsed.validation.summary.errors == 0

            record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-REVERSED")
            await _add_job_row(session, record)
            await facts.ingest_report_batch(
                session,
                record=record,
                normalized=parsed.normalized,
            )

            with pytest.raises(HTTPException) as exc_info:
                await facts.build_dashboard_for_range(
                    session,
                    start_date=date(2026, 5, 1),
                    end_date=date(2026, 4, 30),
                )

            assert exc_info.value.status_code == 400
        finally:
            await tx.rollback()
            await engine.dispose()


@pytest.mark.asyncio
async def test_cumulative_fact_dashboard_reports_no_data_for_empty_date() -> None:
    async with SessionLocal() as session:
        tx = await session.begin()
        try:
            parsed = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-NO-DATA"))
            assert parsed.validation.summary.errors == 0

            record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-NO-DATA")
            await _add_job_row(session, record)
            await facts.ingest_report_batch(
                session,
                record=record,
                normalized=parsed.normalized,
            )

            with pytest.raises(HTTPException) as exc_info:
                await facts.build_dashboard_for_range(
                    session,
                    start_date=date(2099, 1, 1),
                    end_date=date(2099, 1, 1),
                )

            assert exc_info.value.status_code == 404
            assert exc_info.value.detail == "Aucune donnée disponible pour cette période."
        finally:
            await tx.rollback()
            await engine.dispose()


@pytest.mark.asyncio
async def test_cumulative_fact_dashboard_does_not_use_future_snapshot() -> None:
    async with SessionLocal() as session:
        tx = await session.begin()
        try:
            parsed = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-MARCH"))
            assert parsed.validation.summary.errors == 0

            record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-MARCH")
            await _add_job_row(session, record)
            await facts.ingest_report_batch(
                session,
                record=record,
                normalized=parsed.normalized,
            )

            dashboard = await facts.build_dashboard_for_range(
                session,
                start_date=date(2026, 4, 1),
                end_date=date(2026, 4, 9),
            )

            assert dashboard["cards"]["inscriptions"]["rsuMenagesTotal"]["raw"] is None
            assert dashboard["charts"]["asdEntreesSorties"]["series"][0]["month"] == "2026-04"
        finally:
            await tx.rollback()
            await engine.dispose()


@pytest.mark.asyncio
async def test_duplicate_replace_supersedes_active_batch_without_deleting_history() -> None:
    async with SessionLocal() as session:
        tx = await session.begin()
        try:
            original = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-REPLACE"))
            replacement = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-REPLACE"))
            assert original.validation.summary.errors == 0
            assert replacement.validation.summary.errors == 0

            original_record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-REPLACE")
            replacement_record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-REPLACE")
            await _add_job_row(session, original_record)
            await _add_job_row(session, replacement_record)
            old_batch = await facts.ingest_report_batch(
                session,
                record=original_record,
                normalized=original.normalized,
            )
            new_batch = await facts.ingest_report_batch(
                session,
                record=replacement_record,
                normalized=replacement.normalized,
            )

            rows = list(
                await session.scalars(
                    select(ReportUploadBatch).where(
                        ReportUploadBatch.id_chargement == "LOAD-REPLACE"
                    )
                )
            )
            assert len(rows) == 2
            assert {row.is_active for row in rows} == {False, True}
            assert old_batch.is_active is False
            assert old_batch.superseded_by_batch_id == new_batch.id
            assert new_batch.is_active is True

            dashboard = await facts.build_dashboard_for_range(
                session,
                start_date=date(2026, 4, 1),
                end_date=date(2026, 4, 30),
            )
            assert dashboard["cards"]["inscriptions"]["rsuMenagesTotal"]["raw"] == 110
        finally:
            await tx.rollback()
            await engine.dispose()


@pytest.mark.asyncio
async def test_overlapping_period_upload_is_rejected() -> None:
    async with SessionLocal() as session:
        tx = await session.begin()
        try:
            original = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-OVERLAP-1"))
            overlap = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-OVERLAP-2"))
            assert original.validation.summary.errors == 0
            assert overlap.validation.summary.errors == 0
            overlap.normalized["metadata"]["debut_periode"] = "2026-04-15"
            overlap.normalized["metadata"]["fin_periode"] = "2026-05-15"

            original_record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-OVERLAP-1")
            overlap_record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-OVERLAP-2")
            await _add_job_row(session, original_record)
            await _add_job_row(session, overlap_record)
            await facts.ingest_report_batch(
                session,
                record=original_record,
                normalized=original.normalized,
            )

            with pytest.raises(ValueError, match="chevauchante"):
                await facts.ingest_report_batch(
                    session,
                    record=overlap_record,
                    normalized=overlap.normalized,
                )
        finally:
            await tx.rollback()
            await engine.dispose()


@pytest.mark.asyncio
async def test_overlapping_period_replace_supersedes_active_batches() -> None:
    async with SessionLocal() as session:
        tx = await session.begin()
        try:
            original = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-OVERLAP-A"))
            overlap = parser.parse_workbook(make_report_workbook(id_chargement="LOAD-OVERLAP-B"))
            assert original.validation.summary.errors == 0
            assert overlap.validation.summary.errors == 0
            overlap.normalized["metadata"]["debut_periode"] = "2026-04-15"
            overlap.normalized["metadata"]["fin_periode"] = "2026-05-15"

            original_record = _job_record(str(uuid.uuid4()), id_chargement="LOAD-OVERLAP-A")
            overlap_record = _job_record(
                str(uuid.uuid4()),
                id_chargement="LOAD-OVERLAP-B",
                replace=True,
            )
            await _add_job_row(session, original_record)
            await _add_job_row(session, overlap_record)
            old_batch = await facts.ingest_report_batch(
                session,
                record=original_record,
                normalized=original.normalized,
            )
            new_batch = await facts.ingest_report_batch(
                session,
                record=overlap_record,
                normalized=overlap.normalized,
            )

            assert old_batch.is_active is False
            assert old_batch.status == "superseded"
            assert old_batch.superseded_by_batch_id == new_batch.id
            assert new_batch.is_active is True
        finally:
            await tx.rollback()
            await engine.dispose()


def test_local_storage_fallback_roundtrip(tmp_path: Path) -> None:
    storage = LocalObjectStorage(tmp_path)
    storage.put_json("rsu-dashboard/jobs/job-1/validation.json", {"ok": True})
    assert storage.get_json("rsu-dashboard/jobs/job-1/validation.json") == {"ok": True}


@pytest.mark.asyncio
async def test_async_job_lifecycle_local_repository(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "report_job_repository", "local")
    monkeypatch.setattr(settings, "storage_driver", "local")
    monkeypatch.setattr(settings, "storage_local_path", tmp_path)
    monkeypatch.setattr(
        service,
        "_LOCAL_REPOSITORY",
        LocalJsonReportJobRepository(tmp_path / "jobs.json"),
    )
    get_storage.cache_clear()
    content = make_report_workbook(id_chargement="LOAD-JOB")
    upload = UploadFile(file=io.BytesIO(content), filename="rapport.xlsx")
    response = await service.create_report_job(
        None,  # type: ignore[arg-type]
        upload=upload,
        replace=False,
        created_by=None,
    )
    assert response.status == "queued"
    status_before = await service.get_job_status(service._LOCAL_REPOSITORY, response.job_id)
    assert status_before.status == "queued"
    assert await service.process_next_report_job() is True
    status_after = await service.get_job_status(service._LOCAL_REPOSITORY, response.job_id)
    assert status_after.status == "succeeded"
    latest = await service.get_latest_job_status(service._LOCAL_REPOSITORY)
    assert latest.job_id == response.job_id
    dashboard = await service.get_dashboard_json(service._LOCAL_REPOSITORY, response.job_id)
    assert dashboard["meta"]["idChargement"] == "LOAD-JOB"
    get_storage.cache_clear()


@pytest.mark.asyncio
async def test_async_job_generates_upload_id_when_workbook_omits_id_chargement(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "report_job_repository", "local")
    monkeypatch.setattr(settings, "storage_driver", "local")
    monkeypatch.setattr(settings, "storage_local_path", tmp_path)
    monkeypatch.setattr(
        service,
        "_LOCAL_REPOSITORY",
        LocalJsonReportJobRepository(tmp_path / "jobs.json"),
    )
    get_storage.cache_clear()

    content = make_report_workbook(include_id_chargement=False)
    upload = UploadFile(file=io.BytesIO(content), filename="rapport.xlsx")
    response = await service.create_report_job(
        None,  # type: ignore[arg-type]
        upload=upload,
        replace=False,
        created_by=None,
    )

    record = await service._LOCAL_REPOSITORY.get(response.job_id)
    assert record is not None
    assert record.id_chargement == f"JOB_{response.job_id}"

    assert await service.process_next_report_job() is True
    dashboard = await service.get_dashboard_json(service._LOCAL_REPOSITORY, response.job_id)
    assert dashboard["meta"]["idChargement"] == f"JOB_{response.job_id}"
    get_storage.cache_clear()
