from __future__ import annotations

import io
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from reports.schemas import ValidationMessage, ValidationResult, ValidationSummary

MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")

SYSTEM_METADATA_DEFAULTS: dict[str, Any] = {
    "code_langue": "fr",
    "indicateur_mois_partiel": False,
    "suffixe_libelle_mois_partiel": "",
    "type_unite_graphique_rsu": "MENAGES",
    "methode_tendance": "REGRESSION_LINEAIRE",
    "utiliser_regles_montant_secours": True,
    "charge_par": "system",
    "commentaires": "",
}

SYSTEM_CODE_INDEX: dict[str, dict[str, str]] = {
    "code_registre": {"rnp": "RNP", "rsu": "RSU"},
    "type_unite": {"menages": "MENAGES", "personnes": "PERSONNES"},
    "mode_source": {
        "saisie": "SAISIE",
        "diff_snapshot_cumule": "DIFF_SNAPSHOT_CUMULE",
    },
    "methode_tendance": {
        "regression_lineaire": "REGRESSION_LINEAIRE",
        "moyenne_mobile": "MOYENNE_MOBILE",
        "aucune": "AUCUNE",
    },
    "code_perimetre_programme": {
        "asd": "ASD",
        "amo_tadamon": "AMO_TADAMON",
    },
    "code_programme": {
        "asd": "ASD",
        "amo_tadamon": "AMO_TADAMON",
    },
    "unite_beneficiaire": {"menages": "MENAGES"},
    "type_montant": {"standard": "STANDARD"},
}

SYSTEM_MANAGED_OPEN_CODE_FIELDS = {
    "code_type_famille",
    "code_niveau_risque",
    "code_motif_blocage",
}


@dataclass(frozen=True)
class SectionSpec:
    sheet: str
    key: str
    label: str
    columns: tuple[str, ...]
    aliases: dict[str, tuple[str, ...]] = field(default_factory=dict)
    optional_columns: tuple[str, ...] = ()
    sheet_aliases: tuple[str, ...] = ()
    required: bool = True

    def alias_index(self) -> dict[str, str]:
        out: dict[str, str] = {}
        for column in self.columns:
            out[_normalize_key(column)] = column
            for alias in self.aliases.get(column, ()):
                out[_normalize_key(alias)] = column
        return out

    def required_columns(self) -> tuple[str, ...]:
        optional = set(self.optional_columns)
        return tuple(column for column in self.columns if column not in optional)


PARAMETERS = SectionSpec(
    sheet="01_Parametres",
    key="parameters",
    label="Paramètres",
    columns=(
        "id_chargement",
        "version_fichier",
        "date_rapport",
        "debut_periode",
        "fin_periode",
        "date_reference_donnees",
        "date_heure_extraction",
        "systeme_source",
        "code_langue",
        "mois_reporting_courant",
        "indicateur_mois_partiel",
        "suffixe_libelle_mois_partiel",
        "type_unite_graphique_rsu",
        "methode_tendance",
        "utiliser_regles_montant_secours",
        "charge_par",
        "commentaires",
    ),
    optional_columns=(
        "id_chargement",
        "code_langue",
        "mois_reporting_courant",
        "indicateur_mois_partiel",
        "suffixe_libelle_mois_partiel",
        "type_unite_graphique_rsu",
        "methode_tendance",
        "utiliser_regles_montant_secours",
        "charge_par",
        "commentaires",
    ),
)

REGIONS = SectionSpec(
    sheet="02_Regions",
    key="regions",
    label="Régions",
    columns=("code_region", "nom_region", "ordre_affichage"),
    aliases={
        "nom_region": ("region", "libelle_region", "libelle"),
        "ordre_affichage": ("ordre", "ordre_region"),
    },
    required=False,
)

PROVINCES = SectionSpec(
    sheet="03_Provinces",
    key="provinces",
    label="Provinces",
    columns=("code_region", "nom_province", "ordre_affichage"),
    aliases={
        "nom_province": ("province", "libelle_province", "libelle"),
        "ordre_affichage": ("ordre", "ordre_province"),
    },
    optional_columns=("ordre_affichage",),
    required=False,
)

CODES = SectionSpec(
    sheet="04_Codes",
    key="codes",
    label="Codes",
    columns=("type_code", "code", "libelle", "ordre_affichage"),
    aliases={
        "type_code": ("domaine", "code_liste", "categorie", "type", "nom_champ"),
        "code": ("code_valeur", "valeur", "code_value"),
        "libelle": ("label", "libelle_fr", "nom", "description"),
        "ordre_affichage": ("ordre", "ordre_tri"),
    },
    optional_columns=("ordre_affichage",),
    required=False,
)

RSU_STOCK = SectionSpec(
    sheet="10_RSU",
    key="rsu_stock",
    label="Stock RNP/RSU",
    columns=(
        "id_chargement",
        "date_reference",
        "code_registre",
        "type_unite",
        "total_cumule",
        "systeme_source",
        "commentaires",
    ),
    optional_columns=("id_chargement",),
    sheet_aliases=("10_RSU_Stock",),
)

RSU_FLOW = SectionSpec(
    sheet="10_RSU",
    key="rsu_new_registrations",
    label="Nouvelles inscriptions RSU",
    columns=(
        "id_chargement",
        "debut_periode",
        "fin_periode",
        "date_evenement",
        "mois_evenement",
        "code_region",
        "nom_province",
        "type_unite",
        "nb_nouvelles_inscriptions",
        "mode_source",
        "systeme_source",
        "commentaires",
    ),
    optional_columns=("id_chargement",),
    sheet_aliases=("11_RSU_Nouvelles_Inscriptions",),
)

RSU_ANNOTATIONS = SectionSpec(
    sheet="10_RSU",
    key="rsu_annotations",
    label="Annotations RSU",
    columns=("code_graphique", "mois_evenement", "libelle_annotation", "commentaires"),
    sheet_aliases=("12_RSU_Annotations",),
    required=False,
)

PROGRAM_STOCK_COLUMNS = (
    "id_chargement",
    "date_reference",
    "type_unite",
    "nb_actifs",
    "systeme_source",
    "commentaires",
)
PROGRAM_FLOW_COLUMNS = (
    "id_chargement",
    "debut_periode",
    "fin_periode",
    "date_evenement",
    "mois_evenement",
    "code_region",
    "nom_province",
    "nb_entrants_menages",
    "nb_sortants_menages",
    "nb_entrants_personnes",
    "nb_sortants_personnes",
    "montant_mensuel_entrants_dh",
    "montant_mensuel_sortants_dh",
    "systeme_source",
    "commentaires",
)
RESCORING_COLUMNS = (
    "id_chargement",
    "debut_periode",
    "fin_periode",
    "date_evenement",
    "mois_evenement",
    "code_region",
    "nom_province",
    "nb_sortants_menages",
    "nb_sortants_personnes",
    "nb_menages_non_communiques",
    "montant_mensuel_arrete_dh",
    "systeme_source",
    "commentaires",
)
FRAUD_COLUMNS = (
    "id_chargement",
    "debut_periode",
    "fin_periode",
    "date_evenement",
    "mois_evenement",
    "code_region",
    "nom_province",
    "radiated_hh_count",
    "radiated_persons",
    "montant_mensuel_arrete_dh",
    "systeme_source",
    "commentaires",
)
FRAUD_ALIASES: dict[str, tuple[str, ...]] = {
    "radiated_hh_count": ("nb_menages_radies",),
    "radiated_persons": ("nb_personnes_radiees",),
}

ASD_STOCK = SectionSpec(
    "20_ASD",
    "asd_stock",
    "Stock ASD",
    PROGRAM_STOCK_COLUMNS,
    optional_columns=("id_chargement",),
    sheet_aliases=("20_ASD_Stock",),
)
ASD_FLOW = SectionSpec(
    "20_ASD",
    "asd_flow",
    "Flux entrants/sortants ASD",
    PROGRAM_FLOW_COLUMNS,
    optional_columns=("id_chargement",),
    sheet_aliases=("21_ASD_Flux",),
)
ASD_RESCORING = SectionSpec(
    "20_ASD",
    "asd_rescoring",
    "Rescoring ASD",
    RESCORING_COLUMNS,
    optional_columns=("id_chargement",),
    sheet_aliases=("22_ASD_Rescoring",),
)
ASD_FRAUD = SectionSpec(
    "20_ASD",
    "asd_fraud",
    "Radiation pour fraude ASD",
    FRAUD_COLUMNS,
    FRAUD_ALIASES,
    optional_columns=("id_chargement",),
    sheet_aliases=("23_ASD_Fraude",),
)

AMO_STOCK = SectionSpec(
    "30_AMO_Tadamon",
    "amo_stock",
    "Stock AMO Tadamon",
    PROGRAM_STOCK_COLUMNS,
    optional_columns=("id_chargement",),
    sheet_aliases=("30_AMO_Tadamon_Stock",),
)
AMO_FLOW = SectionSpec(
    "30_AMO_Tadamon",
    "amo_flow",
    "Flux entrants/sortants AMO Tadamon",
    PROGRAM_FLOW_COLUMNS,
    optional_columns=("id_chargement",),
    sheet_aliases=("31_AMO_Tadamon_Flux",),
)
AMO_RESCORING = SectionSpec(
    "30_AMO_Tadamon",
    "amo_rescoring",
    "Rescoring AMO Tadamon",
    RESCORING_COLUMNS,
    optional_columns=("id_chargement",),
    sheet_aliases=("32_AMO_Tadamon_Rescoring",),
)
AMO_FRAUD = SectionSpec(
    "30_AMO_Tadamon",
    "amo_fraud",
    "Radiation pour fraude AMO Tadamon",
    FRAUD_COLUMNS,
    FRAUD_ALIASES,
    optional_columns=("id_chargement",),
    sheet_aliases=("33_AMO_Tadamon_Fraude",),
)

FMS_TREATMENT = SectionSpec(
    sheet="40_FMS",
    key="fms_treatment",
    label="Traitement FMS",
    columns=(
        "id_chargement",
        "date_reference",
        "code_type_famille",
        "code_niveau_risque",
        "code_perimetre_programme",
        "code_region",
        "nom_province",
        "demandes_injectees",
        "demandes_traitees",
        "doute_confirme",
        "doute_leve",
        "en_attente",
        "systeme_source",
        "commentaires",
    ),
    optional_columns=("id_chargement",),
    sheet_aliases=("40_FMS_Traitement",),
)

FMS_BLOCKED = SectionSpec(
    sheet="40_FMS",
    key="fms_blocked",
    label="Ménages bloqués",
    columns=(
        "id_chargement",
        "date_reference",
        "code_motif_blocage",
        "code_type_famille",
        "code_region",
        "nom_province",
        "nb_menages_bloques",
        "nb_personnes_bloquees",
        "systeme_source",
        "commentaires",
    ),
    optional_columns=("id_chargement",),
    sheet_aliases=("41_FMS_Menages_Bloques",),
)

AMOUNT_RULES = SectionSpec(
    sheet="50_Regles_Montant",
    key="amount_rules",
    label="Règles de montant",
    columns=(
        "code_programme",
        "date_effet_debut",
        "date_effet_fin",
        "code_type_famille",
        "unite_beneficiaire",
        "type_montant",
        "montant_mensuel_dh",
        "facteur_annualisation",
        "commentaires",
    ),
    required=False,
)

SECTION_SPECS = (
    PARAMETERS,
    REGIONS,
    PROVINCES,
    CODES,
    RSU_STOCK,
    RSU_FLOW,
    RSU_ANNOTATIONS,
    ASD_STOCK,
    ASD_FLOW,
    ASD_RESCORING,
    ASD_FRAUD,
    AMO_STOCK,
    AMO_FLOW,
    AMO_RESCORING,
    AMO_FRAUD,
    FMS_TREATMENT,
    FMS_BLOCKED,
    AMOUNT_RULES,
)

DATE_FIELDS = {
    "date_rapport",
    "debut_periode",
    "fin_periode",
    "date_reference_donnees",
    "date_reference",
    "date_evenement",
    "date_effet_debut",
    "date_effet_fin",
}
DATETIME_FIELDS = {"date_heure_extraction"}
MONTH_FIELDS = {"mois_evenement", "mois_reporting_courant"}
INTEGER_FIELDS = {
    "total_cumule",
    "nb_nouvelles_inscriptions",
    "nb_actifs",
    "nb_entrants_menages",
    "nb_sortants_menages",
    "nb_entrants_personnes",
    "nb_sortants_personnes",
    "nb_menages_non_communiques",
    "radiated_hh_count",
    "radiated_persons",
    "demandes_injectees",
    "demandes_traitees",
    "doute_confirme",
    "doute_leve",
    "en_attente",
    "nb_menages_bloques",
    "nb_personnes_bloquees",
    "ordre_affichage",
}
AMOUNT_FIELDS = {
    "montant_mensuel_entrants_dh",
    "montant_mensuel_sortants_dh",
    "montant_mensuel_arrete_dh",
    "montant_mensuel_dh",
    "facteur_annualisation",
}
OPTIONAL_NUMERIC_FIELDS = {"montant_mensuel_arrete_dh", "date_effet_fin", "ordre_affichage"}
CODE_FIELDS = {
    "code_registre",
    "type_unite",
    "mode_source",
    "type_unite_graphique_rsu",
    "methode_tendance",
    "code_type_famille",
    "code_niveau_risque",
    "code_perimetre_programme",
    "code_motif_blocage",
    "code_programme",
    "unite_beneficiaire",
    "type_montant",
}


@dataclass
class RawSection:
    spec: SectionSpec
    header_row: int | None
    column_map: dict[int, str]
    rows: list[dict[str, Any]]


@dataclass
class ParsedWorkbook:
    normalized: dict[str, Any]
    validation: ValidationResult


def parse_workbook(content: bytes, *, job_id: str | None = None) -> ParsedWorkbook:
    messages: list[ValidationMessage] = []
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001
        messages.append(
            ValidationMessage(
                severity="error",
                code="WORKBOOK_INVALIDE",
                message=f"Classeur Excel invalide: {exc}",
            )
        )
        validation = _validation_result(messages, job_id=job_id, id_chargement=None)
        return ParsedWorkbook(normalized={}, validation=validation)

    sheet_map = {name.strip(): name for name in workbook.sheetnames}

    raw_sections: dict[str, RawSection] = {}
    specs_by_sheet: dict[str, list[SectionSpec]] = {}
    for spec in SECTION_SPECS:
        resolved_sheet = _resolve_sheet(spec, sheet_map)
        if resolved_sheet is None:
            if spec.required:
                expected = " ou ".join((*spec.sheet_aliases, spec.sheet))
                messages.append(
                    ValidationMessage(
                        severity="error",
                        sheet=spec.sheet_aliases[0] if spec.sheet_aliases else spec.sheet,
                        section=spec.label,
                        code="FEUILLE_OU_SECTION_OBLIGATOIRE_MANQUANTE",
                        message=(
                            f"Feuille ou section obligatoire manquante: "
                            f"{spec.label} ({expected})"
                        ),
                    )
                )
            continue
        specs_by_sheet.setdefault(resolved_sheet, []).append(spec)

    for sheet, specs in specs_by_sheet.items():
        rows = _worksheet_rows(workbook[sheet_map[sheet]])
        found = [_find_header(rows, spec, messages) for spec in specs]
        header_rows = [
            section.header_row
            for section in found
            if section is not None and section.header_row is not None
        ]
        for section in found:
            if section is None or section.header_row is None:
                continue
            next_headers = [row for row in header_rows if row > section.header_row]
            end_row = min(next_headers) if next_headers else len(rows) + 1
            section.rows = _parse_section_rows(rows, section, end_row)
            raw_sections[section.spec.key] = section

    normalized = _normalize(raw_sections, messages)
    _validate_references(normalized, messages)
    id_chargement = normalized.get("metadata", {}).get("id_chargement")
    validation = _validation_result(messages, job_id=job_id, id_chargement=id_chargement)
    return ParsedWorkbook(normalized=normalized, validation=validation)


def _resolve_sheet(spec: SectionSpec, sheet_map: dict[str, str]) -> str | None:
    for sheet in spec.sheet_aliases:
        if sheet in sheet_map:
            return sheet
    if spec.sheet in sheet_map:
        return spec.sheet
    return None


def extract_id_chargement(content: bytes) -> str | None:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return None
    if PARAMETERS.sheet not in workbook.sheetnames:
        return None
    rows = _worksheet_rows(workbook[PARAMETERS.sheet])
    messages: list[ValidationMessage] = []
    section = _find_header(rows, PARAMETERS, messages, emit_errors=False)
    if section is None or section.header_row is None:
        return None
    section.rows = _parse_section_rows(rows, section, len(rows) + 1)
    for row in section.rows:
        value = row.get("id_chargement")
        if not _is_blank(value):
            return str(value).strip()
    return None


def _worksheet_rows(worksheet: Any) -> list[list[Any]]:
    return [[cell.value for cell in row] for row in worksheet.iter_rows()]


def _find_header(
    rows: list[list[Any]],
    spec: SectionSpec,
    messages: list[ValidationMessage],
    *,
    emit_errors: bool = True,
) -> RawSection | None:
    alias_index = spec.alias_index()
    best_row: int | None = None
    best_map: dict[int, str] = {}
    best_present: set[str] = set()

    for row_number, row in enumerate(rows, start=1):
        column_map: dict[int, str] = {}
        present: set[str] = set()
        for column_index, value in enumerate(row, start=1):
            key = _normalize_key(value)
            if not key:
                continue
            canonical = alias_index.get(key)
            if canonical:
                column_map[column_index] = canonical
                present.add(canonical)
        if len(present) > len(best_present):
            best_row = row_number
            best_map = column_map
            best_present = present
        if set(spec.required_columns()).issubset(present):
            unknown = _unknown_columns(row, alias_index)
            if emit_errors:
                for column_name in unknown:
                    messages.append(
                        ValidationMessage(
                            severity="warning",
                            sheet=spec.sheet,
                            section=spec.label,
                            row=row_number,
                            column=column_name,
                            code="COLONNE_INCONNUE",
                            message=(
                                f"Colonne inconnue dans la section {spec.label}: "
                                f"{column_name}"
                            ),
                        )
                    )
            return RawSection(spec=spec, header_row=row_number, column_map=column_map, rows=[])

    if not emit_errors:
        return None

    if best_row is None or not best_present:
        if spec.required:
            messages.append(
                ValidationMessage(
                    severity="error",
                    sheet=spec.sheet,
                    section=spec.label,
                    code="SECTION_OBLIGATOIRE_MANQUANTE",
                    message=f"Section obligatoire manquante: {spec.label}",
                )
            )
        return None

    missing = [column for column in spec.required_columns() if column not in best_present]
    for column in missing:
        messages.append(
            ValidationMessage(
                severity="error",
                sheet=spec.sheet,
                section=spec.label,
                row=best_row,
                column=column,
                code="COLONNE_OBLIGATOIRE_MANQUANTE",
                message=(
                    f"Colonne obligatoire manquante dans la section "
                    f"{spec.label}: {column}"
                ),
            )
        )
    return RawSection(spec=spec, header_row=best_row, column_map=best_map, rows=[])


def _parse_section_rows(
    rows: list[list[Any]],
    section: RawSection,
    end_row: int,
) -> list[dict[str, Any]]:
    parsed: list[dict[str, Any]] = []
    if section.header_row is None:
        return parsed
    for row_number in range(section.header_row + 1, end_row):
        row = rows[row_number - 1]
        if sum(1 for value in row if not _is_blank(value)) <= 1:
            continue
        item: dict[str, Any] = {"_row": row_number}
        for column_index, canonical in section.column_map.items():
            item[canonical] = row[column_index - 1] if column_index <= len(row) else None
        if any(not _is_blank(item.get(column)) for column in section.spec.columns):
            parsed.append(item)
    return parsed


def _unknown_columns(row: list[Any], alias_index: dict[str, str]) -> list[str]:
    unknown: list[str] = []
    for value in row:
        raw = "" if value is None else str(value).strip()
        if not raw:
            continue
        if _normalize_key(raw) not in alias_index:
            unknown.append(raw)
    return unknown


def _normalize(
    raw_sections: dict[str, RawSection],
    messages: list[ValidationMessage],
) -> dict[str, Any]:
    normalized: dict[str, Any] = {
        "metadata": {},
        "regions": [],
        "provinces": [],
        "codes": [],
        "rsu_stock": [],
        "rsu_new_registrations": [],
        "rsu_annotations": [],
        "asd_stock": [],
        "asd_flow": [],
        "asd_rescoring": [],
        "asd_fraud": [],
        "amo_stock": [],
        "amo_flow": [],
        "amo_rescoring": [],
        "amo_fraud": [],
        "fms_treatment": [],
        "fms_blocked": [],
        "amount_rules": [],
    }

    for key, section in raw_sections.items():
        rows = [
            _normalize_row(section.spec, row, messages)
            for row in section.rows
        ]
        if key == "parameters":
            if len(rows) != 1:
                messages.append(
                    ValidationMessage(
                        severity="error",
                        sheet=section.spec.sheet,
                        section=section.spec.label,
                        code="PARAMETRES_LIGNE_ACTIVE_INVALIDE",
                        message=(
                            "La feuille 01_Parametres doit contenir exactement "
                            "une ligne active de paramètres."
                        ),
                    )
                )
            if rows:
                normalized["metadata"] = rows[0]
        else:
            normalized[key] = rows

    _apply_metadata_defaults(normalized["metadata"])
    return normalized


def _normalize_row(
    spec: SectionSpec,
    row: dict[str, Any],
    messages: list[ValidationMessage],
) -> dict[str, Any]:
    out: dict[str, Any] = {"_row": row.get("_row")}
    for column in spec.columns:
        value = row.get(column)
        if column in DATE_FIELDS:
            out[column] = _coerce_date(value, spec, row, column, messages)
        elif column in DATETIME_FIELDS:
            out[column] = _coerce_datetime(value, spec, row, column, messages)
        elif column in MONTH_FIELDS:
            out[column] = _coerce_month(value, spec, row, column, messages)
        elif column in INTEGER_FIELDS:
            out[column] = _coerce_number(value, spec, row, column, messages, integer=True)
        elif column in AMOUNT_FIELDS:
            out[column] = _coerce_number(value, spec, row, column, messages, integer=False)
        elif column.startswith("indicateur_") or column.startswith("utiliser_"):
            out[column] = (
                None
                if column in spec.optional_columns and _is_blank(value)
                else _coerce_bool(value)
            )
        else:
            out[column] = None if _is_blank(value) else str(value).strip()
    return out


def _apply_metadata_defaults(metadata: dict[str, Any]) -> None:
    if not metadata:
        return
    for key, default in SYSTEM_METADATA_DEFAULTS.items():
        if _is_blank(metadata.get(key)):
            metadata[key] = default
    if _is_blank(metadata.get("mois_reporting_courant")):
        basis = (
            metadata.get("fin_periode")
            or metadata.get("date_reference_donnees")
            or metadata.get("date_rapport")
        )
        if isinstance(basis, str) and len(basis) >= 7:
            metadata["mois_reporting_courant"] = basis[:7]


def _validate_references(normalized: dict[str, Any], messages: list[ValidationMessage]) -> None:
    metadata = normalized.get("metadata", {})
    id_chargement = metadata.get("id_chargement")

    regions = _region_index(normalized.get("regions", []), messages)
    provinces = _province_index(normalized.get("provinces", []), messages)
    codes = _code_index(normalized.get("codes", []))

    _validate_code_fields(metadata, PARAMETERS, codes, messages)
    for key, spec in _data_section_specs().items():
        for row in normalized.get(key, []):
            _validate_id_chargement(row, spec, id_chargement, messages)
            _validate_code_fields(row, spec, codes, messages)
            _validate_region_province(row, spec, regions, provinces, messages)

    _validate_required_stock_combinations(normalized, messages)
    _validate_fms_coherence(normalized, messages)
    _validate_amount_rule_fallbacks(normalized, messages)


def _data_section_specs() -> dict[str, SectionSpec]:
    return {spec.key: spec for spec in SECTION_SPECS if spec.key != "parameters"}


def _validate_id_chargement(
    row: dict[str, Any],
    spec: SectionSpec,
    expected: str | None,
    messages: list[ValidationMessage],
) -> None:
    if "id_chargement" not in spec.columns or not expected:
        return
    actual = row.get("id_chargement")
    if _is_blank(actual):
        return
    if str(actual) != str(expected):
        messages.append(
            ValidationMessage(
                severity="error",
                sheet=spec.sheet,
                section=spec.label,
                row=row.get("_row"),
                column="id_chargement",
                code="ID_CHARGEMENT_INCOHERENT",
                message=(
                    f"id_chargement incohérent: {actual} "
                    f"(attendu: {expected})"
                ),
            )
        )


def _validate_code_fields(
    row: dict[str, Any],
    spec: SectionSpec,
    codes: dict[str, dict[str, str]],
    messages: list[ValidationMessage],
) -> None:
    for field_name in CODE_FIELDS.intersection(row.keys()):
        value = row.get(field_name)
        if _is_blank(value):
            continue
        type_candidates = _code_type_candidates(field_name)
        accepted_codes: dict[str, str] | None = None
        for candidate in type_candidates:
            if candidate in codes:
                accepted_codes = codes[candidate]
                break
        if accepted_codes is None:
            if field_name in SYSTEM_MANAGED_OPEN_CODE_FIELDS:
                continue
            messages.append(
                ValidationMessage(
                    severity="error",
                    sheet=spec.sheet,
                    section=spec.label,
                    row=row.get("_row"),
                    column=field_name,
                    code="REFERENTIEL_CODE_ABSENT",
                    message=f"Référentiel de codes absent dans le système: {field_name}",
                )
            )
            continue
        if _normalize_key(value) not in accepted_codes:
            messages.append(
                ValidationMessage(
                    severity="error",
                    sheet=spec.sheet,
                    section=spec.label,
                    row=row.get("_row"),
                    column=field_name,
                    code="CODE_INVALIDE",
                    message=f"Code invalide pour {field_name}: {value}",
                )
            )


def _validate_region_province(
    row: dict[str, Any],
    spec: SectionSpec,
    regions: dict[str, dict[str, Any]],
    provinces: dict[str, dict[str, Any]],
    messages: list[ValidationMessage],
) -> None:
    if not regions and not provinces:
        return
    code_region = row.get("code_region")
    if regions and not _is_blank(code_region) and str(code_region) not in regions:
        messages.append(
            ValidationMessage(
                severity="error",
                sheet=spec.sheet,
                section=spec.label,
                row=row.get("_row"),
                column="code_region",
                code="CODE_REGION_INVALIDE",
                message=f"Code région invalide: {code_region}",
            )
        )
    province = row.get("nom_province")
    if _is_blank(province):
        return
    if not provinces:
        return
    province_key = _normalize_key(province)
    if province_key not in provinces:
        messages.append(
            ValidationMessage(
                severity="error",
                sheet=spec.sheet,
                section=spec.label,
                row=row.get("_row"),
                column="nom_province",
                code="PROVINCE_INVALIDE",
                message=f"Province invalide: {province}",
            )
        )
        return
    expected_region = provinces[province_key].get("code_region")
    if (
        not _is_blank(code_region)
        and expected_region
        and str(code_region) != str(expected_region)
    ):
        messages.append(
            ValidationMessage(
                severity="error",
                sheet=spec.sheet,
                section=spec.label,
                row=row.get("_row"),
                column="nom_province",
                code="PROVINCE_REGION_INCOHERENTE",
                message=f"La province {province} n’est pas rattachée à la région {code_region}",
            )
        )


def _validate_required_stock_combinations(
    normalized: dict[str, Any],
    messages: list[ValidationMessage],
) -> None:
    _require_combinations(
        normalized.get("rsu_stock", []),
        RSU_STOCK,
        {("RNP", "PERSONNES"), ("RSU", "MENAGES"), ("RSU", "PERSONNES")},
        messages,
        ("code_registre", "type_unite"),
    )
    _require_combinations(
        normalized.get("asd_stock", []),
        ASD_STOCK,
        {("MENAGES",), ("PERSONNES",)},
        messages,
        ("type_unite",),
        prefix="ASD/",
    )
    _require_combinations(
        normalized.get("amo_stock", []),
        AMO_STOCK,
        {("MENAGES",), ("PERSONNES",)},
        messages,
        ("type_unite",),
        prefix="AMO_TADAMON/",
    )


def _require_combinations(
    rows: list[dict[str, Any]],
    spec: SectionSpec,
    required: set[tuple[str, ...]],
    messages: list[ValidationMessage],
    fields: tuple[str, ...],
    *,
    prefix: str = "",
) -> None:
    actual = {
        tuple(str(row.get(field, "")).upper() for field in fields)
        for row in rows
    }
    for combo in sorted(required - actual):
        label = prefix + "/".join(combo)
        messages.append(
            ValidationMessage(
                severity="error",
                sheet=spec.sheet,
                section=spec.label,
                code="COMBINAISON_STOCK_MANQUANTE",
                message=f"Combinaison de stock obligatoire manquante: {label}",
            )
        )


def _validate_fms_coherence(
    normalized: dict[str, Any],
    messages: list[ValidationMessage],
) -> None:
    for row in normalized.get("fms_treatment", []):
        injected = row.get("demandes_injectees") or 0
        processed = row.get("demandes_traitees") or 0
        confirmed = row.get("doute_confirme") or 0
        cleared = row.get("doute_leve") or 0
        waiting = row.get("en_attente")
        if processed > injected:
            messages.append(
                ValidationMessage(
                    severity="error",
                    sheet=FMS_TREATMENT.sheet,
                    section=FMS_TREATMENT.label,
                    row=row.get("_row"),
                    column="demandes_traitees",
                    code="FMS_TRAITEES_SUPERIEUR_INJECTEES",
                    message="Incohérence FMS: demandes_traitees dépasse demandes_injectees",
                )
            )
        if confirmed + cleared > processed:
            messages.append(
                ValidationMessage(
                    severity="error",
                    sheet=FMS_TREATMENT.sheet,
                    section=FMS_TREATMENT.label,
                    row=row.get("_row"),
                    column="doute_confirme",
                    code="FMS_DOUTES_SUPERIEUR_TRAITEES",
                    message="Incohérence FMS: doute_confirme + doute_leve dépasse demandes_traitees",
                )
            )
        expected_waiting = injected - processed
        tolerance = max(1, math.ceil(injected * 0.01))
        if waiting is not None and abs(waiting - expected_waiting) > tolerance:
            messages.append(
                ValidationMessage(
                    severity="warning",
                    sheet=FMS_TREATMENT.sheet,
                    section=FMS_TREATMENT.label,
                    row=row.get("_row"),
                    column="en_attente",
                    code="FMS_EN_ATTENTE_INCOHERENT",
                    message=(
                        "Incohérence FMS: en_attente ne correspond pas à "
                        "demandes_injectees - demandes_traitees"
                    ),
                )
            )


def _validate_amount_rule_fallbacks(
    normalized: dict[str, Any],
    messages: list[ValidationMessage],
) -> None:
    rules = normalized.get("amount_rules", [])
    if not rules:
        return
    for key, spec, programme in (
        ("asd_fraud", ASD_FRAUD, "ASD"),
        ("amo_fraud", AMO_FRAUD, "AMO_TADAMON"),
        ("asd_rescoring", ASD_RESCORING, "ASD"),
        ("amo_rescoring", AMO_RESCORING, "AMO_TADAMON"),
    ):
        for row in normalized.get(key, []):
            if row.get("montant_mensuel_arrete_dh") is not None:
                continue
            if _find_amount_rule(rules, programme) is None:
                messages.append(
                    ValidationMessage(
                        severity="error",
                        sheet=spec.sheet,
                        section=spec.label,
                        row=row.get("_row"),
                        column="montant_mensuel_arrete_dh",
                        code="REGLE_MONTANT_SECOURS_MANQUANTE",
                        message=(
                            "Montant mensuel arrêté manquant et aucune règle de "
                            f"montant de secours disponible pour {programme}."
                        ),
                    )
                )


def _find_amount_rule(rules: list[dict[str, Any]], programme: str) -> dict[str, Any] | None:
    for rule in rules:
        if str(rule.get("code_programme", "")).upper() == programme:
            return rule
    return None


def _region_index(
    rows: list[dict[str, Any]],
    messages: list[ValidationMessage],
) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = row.get("code_region")
        if _is_blank(code):
            messages.append(
                ValidationMessage(
                    severity="error",
                    sheet=REGIONS.sheet,
                    section=REGIONS.label,
                    row=row.get("_row"),
                    column="code_region",
                    code="CODE_REGION_MANQUANT",
                    message="Code région manquant dans 02_Regions.",
                )
            )
            continue
        index[str(code)] = row
    return index


def _province_index(
    rows: list[dict[str, Any]],
    messages: list[ValidationMessage],
) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    for row in rows:
        province = row.get("nom_province")
        code_region = row.get("code_region")
        if _is_blank(province) or _is_blank(code_region):
            messages.append(
                ValidationMessage(
                    severity="error",
                    sheet=PROVINCES.sheet,
                    section=PROVINCES.label,
                    row=row.get("_row"),
                    code="PROVINCE_REFERENTIEL_INVALIDE",
                    message="Province ou code région manquant dans 03_Provinces.",
                )
            )
            continue
        index[_normalize_key(province)] = row
    return index


def _code_index(rows: list[dict[str, Any]]) -> dict[str, dict[str, str]]:
    index: dict[str, dict[str, str]] = {
        type_code: dict(values) for type_code, values in SYSTEM_CODE_INDEX.items()
    }
    for row in rows:
        type_code = row.get("type_code")
        code = row.get("code")
        if _is_blank(type_code) or _is_blank(code):
            continue
        index.setdefault(_normalize_key(type_code), {})[_normalize_key(code)] = str(
            row.get("libelle") or code
        )
    return index


def _code_type_candidates(field_name: str) -> list[str]:
    normalized = _normalize_key(field_name)
    candidates = [normalized]
    if normalized.startswith("code_"):
        candidates.append(normalized.removeprefix("code_"))
    if normalized == "type_unite_graphique_rsu":
        candidates.append("type_unite")
    if normalized == "code_programme":
        candidates.append("code_perimetre_programme")
    return candidates


def _coerce_date(
    value: Any,
    spec: SectionSpec,
    row: dict[str, Any],
    column: str,
    messages: list[ValidationMessage],
) -> str | None:
    if _is_blank(value):
        if column not in OPTIONAL_NUMERIC_FIELDS and column not in spec.optional_columns:
            messages.append(_invalid_value(spec, row, column, f"Date obligatoire manquante: {column}"))
        return None
    parsed = _parse_date(value)
    if parsed is None:
        messages.append(_invalid_value(spec, row, column, f"Date invalide: {column}"))
        return None
    return parsed.isoformat()


def _coerce_datetime(
    value: Any,
    spec: SectionSpec,
    row: dict[str, Any],
    column: str,
    messages: list[ValidationMessage],
) -> str | None:
    if _is_blank(value):
        messages.append(_invalid_value(spec, row, column, f"Date invalide: {column}"))
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day).isoformat()
    if isinstance(value, int | float):
        try:
            parsed = cast(datetime, from_excel(value))
            return parsed.isoformat()
        except Exception:
            pass
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text).isoformat()
    except ValueError:
        parsed_date = _parse_date(text)
        if parsed_date:
            return datetime(parsed_date.year, parsed_date.month, parsed_date.day).isoformat()
    messages.append(_invalid_value(spec, row, column, f"Date invalide: {column}"))
    return None


def _coerce_month(
    value: Any,
    spec: SectionSpec,
    row: dict[str, Any],
    column: str,
    messages: list[ValidationMessage],
) -> str | None:
    if _is_blank(value):
        if column in spec.optional_columns:
            return None
        messages.append(_invalid_value(spec, row, column, f"Mois obligatoire manquant: {column}"))
        return None
    text = str(value).strip()
    if MONTH_RE.match(text):
        return text
    messages.append(
        _invalid_value(
            spec,
            row,
            column,
            f"Format de mois invalide pour {column}: attendu AAAA-MM",
        )
    )
    return None


def _coerce_number(
    value: Any,
    spec: SectionSpec,
    row: dict[str, Any],
    column: str,
    messages: list[ValidationMessage],
    *,
    integer: bool,
) -> int | float | None:
    if _is_blank(value):
        if column in OPTIONAL_NUMERIC_FIELDS or column in spec.optional_columns:
            return None
        messages.append(_invalid_value(spec, row, column, f"Valeur numérique invalide: {column}"))
        return 0 if integer else 0.0
    try:
        number = float(str(value).replace(",", ".")) if isinstance(value, str) else float(value)
    except (TypeError, ValueError):
        messages.append(_invalid_value(spec, row, column, f"Valeur numérique invalide: {column}"))
        return 0 if integer else 0.0
    if number < 0:
        messages.append(_invalid_value(spec, row, column, f"Valeur numérique négative: {column}"))
    return int(number) if integer else number


def _invalid_value(
    spec: SectionSpec,
    row: dict[str, Any],
    column: str,
    message: str,
) -> ValidationMessage:
    return ValidationMessage(
        severity="error",
        sheet=spec.sheet,
        section=spec.label,
        row=row.get("_row"),
        column=column,
        code="VALEUR_INVALIDE",
        message=message,
    )


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float):
        try:
            parsed = cast(datetime, from_excel(value))
            return parsed.date()
        except Exception:
            return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def _coerce_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, int | float):
        return bool(value)
    text = str(value).strip().lower()
    return text in {"true", "vrai", "oui", "1", "yes", "y"}


def _validation_result(
    messages: list[ValidationMessage],
    *,
    job_id: str | None,
    id_chargement: str | None,
) -> ValidationResult:
    summary = ValidationSummary(
        errors=sum(1 for message in messages if message.severity == "error"),
        warnings=sum(1 for message in messages if message.severity == "warning"),
        infos=sum(1 for message in messages if message.severity == "info"),
    )
    return ValidationResult(
        job_id=job_id,
        id_chargement=id_chargement,
        summary=summary,
        messages=messages,
    )


def _normalize_key(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and not value.strip()
