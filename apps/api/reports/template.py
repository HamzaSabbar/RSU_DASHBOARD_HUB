from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True)
class TemplateSheet:
    name: str
    headers: tuple[str, ...]
    required: tuple[str, ...]
    note: str


DATE_COMMENT = "Format accepté: YYYY-MM-DD, DD/MM/YYYY, DD-MM-YYYY ou date Excel."
GEOGRAPHY_COMMENT = (
    "Nom métier, pas de code interne. Exemple: casablanca-settat, rabat, kenitra."
)
INTEGER_COMMENT = "Nombre entier positif ou zéro."
AMOUNT_COMMENT = "Montant en dirhams. Laisser vide si la source ne le fournit pas."


TEMPLATE_SHEETS: tuple[TemplateSheet, ...] = (
    TemplateSheet(
        name="01_Parametres",
        headers=(
            "debut_periode",
            "fin_periode",
            "version_fichier",
            "date_reference_donnees",
            "date_heure_extraction",
            "systeme_source",
            "commentaires",
        ),
        required=("debut_periode", "fin_periode"),
        note="Une seule ligne active décrit la période globale du fichier.",
    ),
    TemplateSheet(
        name="11_RNP_Nouvelles_Inscriptions",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_nouvelles_inscriptions",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_nouvelles_inscriptions",
        ),
        note="Flux des nouvelles inscriptions individuelles au RNP.",
    ),
    TemplateSheet(
        name="12_RSU_Nouvelles_Inscriptions",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_nouveaux_menages_rsu",
            "nb_nouvelles_personnes_rsu",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_nouveaux_menages_rsu",
            "nb_nouvelles_personnes_rsu",
        ),
        note="Flux quotidien des nouveaux ménages RSU et des personnes couvertes.",
    ),
    TemplateSheet(
        name="21_ASD_Flux",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_entrants_menages",
            "nb_entrants_personnes",
            "nb_sortants_menages",
            "nb_sortants_personnes",
            "montant_mensuel_entrants_dh",
            "montant_mensuel_sortants_dh",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_entrants_menages",
            "nb_entrants_personnes",
            "nb_sortants_menages",
            "nb_sortants_personnes",
        ),
        note="Flux quotidien entrants et sortants ASD par province.",
    ),
    TemplateSheet(
        name="22_ASD_Rescoring",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_sortants_menages",
            "nb_menages_non_communiques",
            "montant_mensuel_arrete_dh",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_sortants_menages",
            "nb_menages_non_communiques",
        ),
        note="Sorties ASD liées au rescoring.",
    ),
    TemplateSheet(
        name="23_ASD_Fraude",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "radiated_hh_count",
            "montant_mensuel_arrete_dh",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "radiated_hh_count",
        ),
        note="Radiations ASD pour fraude.",
    ),
    TemplateSheet(
        name="31_AMO_Tadamon_Flux",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_entrants_menages",
            "nb_entrants_personnes",
            "nb_sortants_menages",
            "nb_sortants_personnes",
            "montant_mensuel_entrants_dh",
            "montant_mensuel_sortants_dh",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_entrants_menages",
            "nb_entrants_personnes",
            "nb_sortants_menages",
            "nb_sortants_personnes",
        ),
        note="Flux quotidien entrants et sortants AMO Tadamon par province.",
    ),
    TemplateSheet(
        name="32_AMO_Tadamon_Rescoring",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_sortants_menages",
            "nb_menages_non_communiques",
            "montant_mensuel_arrete_dh",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "nb_sortants_menages",
            "nb_menages_non_communiques",
        ),
        note="Sorties AMO Tadamon liées au rescoring.",
    ),
    TemplateSheet(
        name="33_AMO_Tadamon_Fraude",
        headers=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "radiated_hh_count",
            "montant_mensuel_arrete_dh",
        ),
        required=(
            "date_evenement",
            "nom_region",
            "nom_province",
            "radiated_hh_count",
        ),
        note="Radiations AMO Tadamon pour fraude.",
    ),
    TemplateSheet(
        name="40_FMS_Traitement",
        headers=(
            "date_evenement",
            "type_famille",
            "niveau_risque",
            "nom_region",
            "nom_province",
            "demandes_injectees",
            "demandes_traitees",
            "doute_confirme",
            "doute_leve",
            "en_attente",
        ),
        required=(
            "date_evenement",
            "type_famille",
            "niveau_risque",
            "nom_region",
            "nom_province",
            "demandes_injectees",
            "demandes_traitees",
            "doute_confirme",
            "doute_leve",
            "en_attente",
        ),
        note="Traitements FMS par famille, risque, région et province.",
    ),
    TemplateSheet(
        name="41_FMS_Menages_Bloques",
        headers=(
            "date_evenement",
            "motif_blocage",
            "type_famille",
            "nom_region",
            "nom_province",
            "nb_menages_bloques",
        ),
        required=(
            "date_evenement",
            "motif_blocage",
            "type_famille",
            "nom_region",
            "nom_province",
            "nb_menages_bloques",
        ),
        note="Ménages bloqués par motif, famille, région et province.",
    ),
)


def build_excel_template() -> bytes:
    wb = Workbook()
    guide = wb.active
    guide.title = "00_Guide"
    _build_guide_sheet(guide)
    for spec in TEMPLATE_SHEETS:
        _build_data_sheet(wb.create_sheet(spec.name), spec)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_guide_sheet(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Modèle de chargement RSU Dashboard"
    ws["A1"].font = Font(bold=True, size=16, color="14140F")
    ws["A3"] = "Règles générales"
    ws["A3"].font = Font(bold=True, color="0A3D2A")
    rules = (
        "Remplir uniquement les feuilles listées dans ce fichier.",
        "Ne pas ajouter id_chargement, code_region ou mois_evenement.",
        "Chaque ligne de fait doit contenir date_evenement, au niveau quotidien.",
        "Les exemples ci-dessous sont indicatifs et ne sont pas des lignes à importer.",
        "Format de date recommandé: YYYY-MM-DD.",
    )
    for row_index, rule in enumerate(rules, start=4):
        ws.cell(row_index, 1).value = f"- {rule}"

    start_row = 11
    ws.cell(start_row, 1).value = "Feuille"
    ws.cell(start_row, 2).value = "Description"
    ws.cell(start_row, 3).value = "Champs obligatoires"
    for cell in ws[start_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A3D2A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for offset, spec in enumerate(TEMPLATE_SHEETS, start=1):
        row = start_row + offset
        ws.cell(row, 1).value = spec.name
        ws.cell(row, 2).value = spec.note
        ws.cell(row, 3).value = ", ".join(spec.required)
        for col in range(1, 4):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)

    ws.cell(start_row + len(TEMPLATE_SHEETS) + 3, 1).value = "Exemples de valeurs"
    ws.cell(start_row + len(TEMPLATE_SHEETS) + 3, 1).font = Font(bold=True, color="0A3D2A")
    examples = (
        ("nom_region", "casablanca-settat, rabat-sale-kenitra"),
        ("nom_province", "casablanca, rabat, kenitra"),
        (
            "type_famille",
            "Individuels, Mariés avec enfants, Mono-parental, Multi-noyaux",
        ),
        ("niveau_risque", "eleve, moyen, faible"),
        ("motif_blocage", "fms, multi, indiv"),
    )
    for row_index, (field, value) in enumerate(examples, start=start_row + len(TEMPLATE_SHEETS) + 4):
        ws.cell(row_index, 1).value = field
        ws.cell(row_index, 2).value = value

    widths = (34, 62, 92)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _build_data_sheet(ws: Worksheet, spec: TemplateSheet) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    required = set(spec.required)
    for column_index, header in enumerate(spec.headers, start=1):
        cell = ws.cell(1, column_index)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            "solid",
            fgColor="0A3D2A" if header in required else "4A4A44",
        )
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="DCDCDC"))
        cell.comment = Comment(_header_comment(header, header in required), "RSU Dashboard")
        ws.column_dimensions[get_column_letter(column_index)].width = _column_width(header)

    ws["A3"] = spec.note
    ws["A3"].font = Font(italic=True, color="707070")
    ws["A4"] = "Les lignes de données doivent commencer à la ligne 2, directement sous l'en-tête."
    ws["A4"].font = Font(italic=True, color="707070")


def _header_comment(header: str, required: bool) -> str:
    parts: list[str] = ["Obligatoire." if required else "Optionnel."]
    if header in {"debut_periode", "fin_periode", "date_evenement", "date_reference_donnees"}:
        parts.append(DATE_COMMENT)
    elif header == "date_heure_extraction":
        parts.append("Date/heure d'extraction, par exemple 2026-04-30T08:30:00.")
    elif header in {"nom_region", "nom_province"}:
        parts.append(GEOGRAPHY_COMMENT)
    elif header in {"type_famille", "niveau_risque", "motif_blocage", "type_unite"}:
        parts.append("Libellé métier source; l'application le normalise.")
    elif _is_amount_header(header):
        parts.append(AMOUNT_COMMENT)
    elif _is_count_header(header):
        parts.append(INTEGER_COMMENT)
    return " ".join(parts)


def _is_count_header(header: str) -> bool:
    return header.startswith(("nb_", "demandes_", "doute_", "en_attente", "radiated_"))


def _is_amount_header(header: str) -> bool:
    return header.startswith("montant_")


def _column_width(header: str) -> int:
    if header in {"commentaires", "date_heure_extraction"}:
        return 28
    if header in {"nom_region", "nom_province"}:
        return 24
    if len(header) > 24:
        return 28
    return max(16, len(header) + 4)
